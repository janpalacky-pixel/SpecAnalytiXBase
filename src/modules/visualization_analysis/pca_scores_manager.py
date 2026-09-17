# src/modules/visualization_analysis/pca_scores_manager.py

import os
import numpy as np
import pandas as pd
from src.modules.utils.app_logger import get_logger
logger = get_logger(__name__)


class PcaScoresManager:
    """Business logic for the PCA / SVD Scores & Loadings dialog.

    Mirrors SVDAnalysisManager's shape (same U/s/Vt/x_axis/explained_variance
    state, same save_results_excel/save_results_text split) since both
    dialogs decompose spectra via the same underlying SVD. The one addition
    here is the pair of derived diagnostics (residual error, Malinowski IND)
    computed once per SVD run and cached, since the Diagnostics tab reads
    them on every redraw rather than just once.
    """

    def __init__(self):
        self.U = None                  # (n_wl x n_comp)  — loadings
        self.s = None                  # (n_comp,)        — singular values
        self.Vt = None                 # (n_comp x n_spec) — scores
        self.x_axis = None
        self.spectrum_labels = None
        self.explained_variance = None  # percent per component
        self.residual_err = None
        self.malinowski_ind = None
        self.last_error = None     # specific reason if compute_svd() returns False
        self.mean_spectrum = None  # per-wavelength mean subtracted before SVD, or None
        self.mean_centered = False  # whether the last compute used mean-centering

    def reset(self):
        """Clear all previous results — same purpose as
        ClusterAnalysisManager.reset()/SVDAnalysisManager's implicit reset
        on recompute: without this, results from a previous, unrelated
        spectra selection could still be read before a fresh computation
        finishes."""
        self.__init__()

    # ------------------------------------------------------------------ #
    # SVD computation                                                      #
    # ------------------------------------------------------------------ #

    def compute_svd(self, spectra, n_components, mean_center=True):
        """Compute SVD from spectra via SVDBackgroundManager, keeping only
        the first n_components. Pure computation — no Qt — so this is safe
        to call from a background QThread (see _ComputeWorker in the
        dialog).

        Args:
            mean_center: if True (the default here — this dialog is framed
                as "PCA", where mean-centering before decomposing is the
                standard convention), subtract the mean spectrum before the
                SVD. See SVDBackgroundManager.compute_svd_from_spectra's
                docstring for the full rationale.

        Returns:
            bool: True if the computation succeeded.
        """
        if not spectra:
            logger.debug("DEBUG: No spectra provided for PCA/SVD computation")
            self.last_error = "No spectra were provided."
            return False

        self.last_error = None

        try:
            from src.modules.data_analysis.svd_background_manager import SVDBackgroundManager
            mgr = SVDBackgroundManager()
            mgr.max_components = n_components
            if not mgr.compute_svd_from_spectra(spectra, mean_center=mean_center):
                # SVDBackgroundManager logs its own reason but doesn't expose
                # a message attribute — the axis-mismatch case is already
                # ruled out by validate_common_x_axis before this dialog
                # ever opens, so a failure reaching here is almost always an
                # unexpected one (e.g. NaN/Inf values, a degenerate matrix).
                # Without setting this, the dialog had nothing to show the
                # user at all on failure — not even a generic message — and
                # simply appeared to do nothing.
                self.last_error = (
                    "SVD computation failed. This can happen if the selected "
                    "spectra contain invalid values (NaN/Inf) or an "
                    "otherwise degenerate data matrix. See the application "
                    "log for the specific reason."
                )
                return False

            n = n_components
            self.U = mgr.U[:, :n]
            self.s = mgr.s[:n]
            self.Vt = mgr.Vt[:n, :]
            self.x_axis = np.asarray(mgr.x_axis, dtype=float)
            self.explained_variance = mgr.explained_variance[:n]
            self.spectrum_labels = [spectrum['label'] for spectrum in spectra]
            self.mean_spectrum = mgr.mean_spectrum
            self.mean_centered = mgr.mean_centered

            # SVD's rank is capped by min(n_wavelengths, n_spectra) — with
            # few spectra selected, the actual decomposition can have fewer
            # components than requested. Re-derive everything below from
            # what was actually computed (len(self.s)) rather than trusting
            # n_components, so nothing downstream reads past the end of
            # these arrays.
            self.residual_err = self.calculate_residual_errors()
            self.malinowski_ind = self.calculate_malinowski_ind()
            return True

        except Exception as e:
            logger.error(f"ERROR: PCA/SVD computation failed: {e}")
            logger.exception("Traceback:")
            self.last_error = f"PCA/SVD computation failed: {e}"
            return False

    def get_n_components(self):
        return 0 if self.U is None else self.U.shape[1]

    def get_n_datapoints(self):
        return 0 if self.U is None else self.U.shape[0]

    def get_n_spectra(self):
        return 0 if self.Vt is None else self.Vt.shape[1]

    # ------------------------------------------------------------------ #
    # Derived diagnostics                                                  #
    # ------------------------------------------------------------------ #

    def calculate_residual_errors(self):
        """Residual standard deviation if only the first m components are
        retained — same formula used by SVD Background's own diagnostics
        and by SVDAnalysisManager. The minimum (or the point where the
        curve flattens) suggests how many components are actually
        significant, with the rest being mostly noise.
        """
        if self.s is None or self.U is None or self.Vt is None:
            return None
        # The true number of spectra, not however many components were
        # retained (self.U.shape[1] is min(n_components requested,
        # n_spectra)) — using that instead of the real spectrum count would
        # make the denominator hit zero for later components, artificially
        # zeroing their residual error instead of computing a real value.
        num_of_spectra = self.Vt.shape[1]
        num_of_spec_points = self.U.shape[0]
        if num_of_spectra == 1:
            return np.array([0.0])

        eigen = self.s ** 2
        num_of_scores = len(self.s)
        E = np.zeros(num_of_scores - 1) if num_of_scores > 1 else np.array([])
        for m in range(num_of_scores - 1):
            expr_1 = np.sum(eigen) - np.sum(eigen[:m + 1])
            expr_2 = (num_of_spec_points - m - 1) * (num_of_spectra - m - 1)
            E[m] = np.sqrt(expr_1 / expr_2) if expr_2 > 0 else 0.0
        return E

    def calculate_malinowski_ind(self):
        """IND(m) = RE(m) / (n_sp - m)^2 — a normalized version of the
        residual error. RE alone decreases monotonically and never has a
        true minimum for spectroscopic data; dividing by (n_sp - m)^2
        creates a competing effect that makes IND turn back upward once m
        exceeds the number of real factors, giving it an actual minimum —
        a more reliable indicator of the optimal component count than RE
        by itself. Most reliable when n_pts >> n_sp; for large or
        nearly-square datasets the minimum can be unreliable (Malinowski,
        E.R. (2002) Factor Analysis in Chemistry, 3rd ed.).
        """
        if self.residual_err is None or self.Vt is None:
            return None
        n_sp = self.Vt.shape[1]
        ind = np.full(len(self.residual_err), np.nan)
        for m in range(len(self.residual_err)):
            denom = (n_sp - (m + 1)) ** 2
            if denom > 0:
                ind[m] = self.residual_err[m] / denom
        return ind

    def hotelling_t2(self, pc_indices):
        """Hotelling's T^2 per spectrum over the given (0-based) component
        indices — sum of squared standardized scores. A standard PCA
        outlier-detection metric: spectra far from the bulk of the data in
        score space (in any of the considered components) get a high T^2.
        """
        if self.Vt is None:
            return None
        scores = self.Vt[pc_indices, :]                 # (k x n_spec)
        std = scores.std(axis=1, ddof=1)
        std[std == 0] = 1.0
        standardized = scores / std[:, None]
        return np.sum(standardized ** 2, axis=0)

    def get_metric_values(self, metric_name):
        """Return the per-component array for any of the seven metrics
        used by the Diagnostics tab, all padded to the same length n
        (Residual error / Malinowski IND are naturally one element
        shorter, so the last component reads NaN — nothing left over to
        measure a residual against). Metric names match the _METRICS
        catalog in the dialog exactly (presentation metadata like colors
        and tooltips stays there; this only needs the name to know which
        array to return)."""
        n = len(self.explained_variance) if self.explained_variance is not None else 0
        if metric_name == 'Singular values (\u03c3)':
            return self.s if self.s is not None else np.full(n, np.nan)
        elif metric_name == 'Eigenvalue (\u03c3\u00b2)':
            return (self.s ** 2) if self.s is not None else np.full(n, np.nan)
        elif metric_name == 'Explained variance (%)':
            return self.explained_variance if self.explained_variance is not None else np.full(n, np.nan)
        elif metric_name == 'Cumulative explained var. (%)':
            return np.cumsum(self.explained_variance) if self.explained_variance is not None else np.full(n, np.nan)
        elif metric_name == 'Cumulative unexplained var. (%)':
            cumvar = np.cumsum(self.explained_variance) if self.explained_variance is not None else np.full(n, np.nan)
            return 100.0 - cumvar
        elif metric_name == 'Residual error E(m)':
            arr = self.residual_err if self.residual_err is not None else np.array([])
            out = np.full(n, np.nan)
            out[:len(arr)] = arr
            return out
        elif metric_name == 'Malinowski IND':
            arr = self.malinowski_ind if self.malinowski_ind is not None else np.array([])
            out = np.full(n, np.nan)
            out[:len(arr)] = arr
            return out
        return np.full(n, np.nan)

    # ------------------------------------------------------------------ #
    # Export                                                               #
    # ------------------------------------------------------------------ #

    def _autofit_excel_columns(self, writer, sheet_name, df):
        """Widen each column enough to show its full header on first open
        — same fix applied to SVD Analysis's Excel export."""
        from openpyxl.utils import get_column_letter
        worksheet = writer.sheets[sheet_name]
        for i, col in enumerate(df.columns):
            width = max(len(str(col)) + 2, 10)
            worksheet.column_dimensions[get_column_letter(i + 1)].width = width

    def _build_export_frames(self, labels, include_scores, include_loadings, include_variance):
        """Shared DataFrame construction for both Excel and text export."""
        n = self.Vt.shape[0]

        scores_df = None
        if include_scores:
            scores_dict = {'Label': labels}
            for i in range(n):
                scores_dict[f'PC{i+1}'] = self.Vt[i, :]
            scores_df = pd.DataFrame(scores_dict)

        loadings_df = None
        if include_loadings:
            loadings_dict = {'x': self.x_axis}
            for i in range(n):
                loadings_dict[f'PC{i+1}'] = self.U[:, i]
            loadings_df = pd.DataFrame(loadings_dict)

        variance_df = None
        if include_variance:
            singular_values = self.s if self.s is not None else np.zeros(n)
            eigenvalues = self.s ** 2 if self.s is not None else np.zeros(n)
            cumvar = np.cumsum(self.explained_variance)
            res_col = np.full(n, np.nan)
            if self.residual_err is not None:
                res_col[:len(self.residual_err)] = self.residual_err
            ind_col = np.full(n, np.nan)
            if self.malinowski_ind is not None:
                ind_col[:len(self.malinowski_ind)] = self.malinowski_ind
            variance_df = pd.DataFrame({
                'Component': [f'PC{i+1}' for i in range(n)],
                'Singular_values': singular_values,
                'Eigenvalue': eigenvalues,
                'Explained_Var_pct': self.explained_variance,
                'Cumulative_Var_pct': cumvar,
                'Residual_Error': res_col,
                'Malinowski_IND': ind_col,
            })

        return scores_df, loadings_df, variance_df

    def save_results_excel(self, filepath, labels, include_scores=True,
                            include_loadings=True, include_variance=True):
        """Save Scores/Loadings/Variance to one Excel workbook with
        separate worksheets. Raises on failure so the controller can
        surface a message; matches ClusterAnalysisManager.save_results_excel
        and SVDAnalysisManager.save_results_excel's contract of returning a
        plain bool for success."""
        if self.Vt is None:
            raise ValueError("No PCA/SVD data available to save")

        scores_df, loadings_df, variance_df = self._build_export_frames(
            labels, include_scores, include_loadings, include_variance)

        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            if scores_df is not None:
                scores_df.to_excel(writer, sheet_name='Scores', index=False,
                                    startrow=1)
                writer.sheets['Scores'].cell(row=1, column=1).value = (
                    "Note: scores are unit-normalized (each PCk column has unit length); they are NOT multiplied by the singular value. Multiply column PCk by the matching value in the Variance sheet Singular_values column to obtain scikit-learn/Jolliffe-convention (sigma-scaled) PCA scores."
                )
                self._autofit_excel_columns(writer, 'Scores', scores_df)
            if loadings_df is not None:
                loadings_df.to_excel(writer, sheet_name='Loadings', index=False)
                self._autofit_excel_columns(writer, 'Loadings', loadings_df)
            if variance_df is not None:
                variance_df.to_excel(writer, sheet_name='Variance', index=False)
                self._autofit_excel_columns(writer, 'Variance', variance_df)
        return True

    def save_results_text(self, save_config):
        """Save Scores/Loadings/Variance to text/CSV file(s) based on
        save_config (file_path, labels, delimiter, precision,
        save_separate, include_scores, include_loadings, include_variance)."""
        if self.Vt is None:
            raise ValueError("No PCA/SVD data available to save")

        file_path = save_config['file_path']
        labels = save_config['labels']
        delimiter = save_config.get('delimiter', '\t')
        precision = save_config.get('precision', 6)
        float_format = f'%.{precision}f'
        include_scores = save_config.get('include_scores', True)
        include_loadings = save_config.get('include_loadings', True)
        include_variance = save_config.get('include_variance', True)

        scores_df, loadings_df, variance_df = self._build_export_frames(
            labels, include_scores, include_loadings, include_variance)

        if save_config.get('save_separate'):
            base = os.path.splitext(file_path)[0]
            if scores_df is not None:
                with open(f'{base}_scores.txt', 'w', newline='') as sf:
                    sf.write('# Note: scores are unit-normalized (each PCk column has unit length); they are NOT multiplied by the singular value. Multiply column PCk by the matching value in the Variance sheet Singular_values column to obtain scikit-learn/Jolliffe-convention (sigma-scaled) PCA scores.\n')
                    scores_df.to_csv(sf, sep=delimiter, index=False, float_format=float_format)
            if loadings_df is not None:
                loadings_df.to_csv(f'{base}_loadings.txt', sep=delimiter, index=False, float_format=float_format)
            if variance_df is not None:
                variance_df.to_csv(f'{base}_variance.txt', sep=delimiter, index=False, float_format=float_format)
        else:
            with open(file_path, 'w', newline='') as f:
                if scores_df is not None:
                    f.write('# Scores\n')
                    f.write('# Note: scores are unit-normalized (each PCk column has unit length); they are NOT multiplied by the singular value. Multiply column PCk by the matching value in the Variance sheet Singular_values column to obtain scikit-learn/Jolliffe-convention (sigma-scaled) PCA scores.\n')
                    scores_df.to_csv(f, sep=delimiter, index=False, float_format=float_format)
                    f.write('\n')
                if loadings_df is not None:
                    f.write('# Loadings\n')
                    loadings_df.to_csv(f, sep=delimiter, index=False, float_format=float_format)
                    f.write('\n')
                if variance_df is not None:
                    f.write('# Variance\n')
                    variance_df.to_csv(f, sep=delimiter, index=False, float_format=float_format)
        return True
