# src/modules/data_io/row_data_converter.py
#
# Row-oriented import for spectra files (spectra in rows).
#
# Expected layout:
#   row 0      -> shared X-scale, one spectral point per column
#   row 1 … N  -> one spectrum's Y-values per row
#
# An optional first COLUMN provides spectrum labels for rows 1..N — this is
# the row-oriented analogue of the header ROW in the standard column format.
#
#   (corner, blank/"x")   x1        x2        x3    …
#   spectrum_A_label      yA1       yA2       yA3   …
#   spectrum_B_label      yB1       yB2       yB3   …
#
# Implementation strategy
# ------------------------
# Row-oriented data is nothing more than the TRANSPOSE of the standard
# column-oriented layout. Rather than duplicate the parsing, detection and
# label-sanitising logic, we tokenize the raw lines, transpose the token
# matrix, rejoin it as tab-separated text, and hand it to the existing,
# well-tested `_parse_column_data()` from table_data_converter. Every
# column-format capability (delimiter/decimal auto-detection reuse, NaN
# handling, label cleanup, zero-padding, Excel support) is inherited for
# free and stays in one place.

import os
from typing import List, Dict, Union, Optional

import numpy as np

from src.modules.utils.app_logger import get_logger
logger = get_logger(__name__)



def _apply_header_row_rows(rows, header_row):
    """Drop everything above *header_row* (0-based). See
    table_data_converter._apply_header_row — same mechanism, and in this
    row-oriented layout the chosen row becomes the shared X-scale row."""
    if header_row is None:
        return rows
    if header_row < 0 or header_row >= len(rows):
        raise ValueError(
            f"Header row {header_row + 1} is out of range — the file only has "
            f"{len(rows)} rows."
        )
    return rows[header_row:]


def _apply_label_column(rows_tokens: List[List[str]], label_column: Optional[int]) -> None:
    """
    Swap the user-chosen label column into position 0, in place.

    Every downstream step (header detection, `_parse_column_data`) always
    treats column 0 as the label source, so moving the real label column
    there — rather than teaching every downstream step about an arbitrary
    index — keeps all of that logic untouched. Column order otherwise
    doesn't matter: each column is evaluated independently based on its
    own first token, not its position.
    """
    if label_column is None or label_column == 0:
        return
    if not rows_tokens or label_column >= len(rows_tokens[0]):
        raise ValueError(
            f"label_column={label_column} is out of range — the file has "
            f"only {len(rows_tokens[0]) if rows_tokens else 0} columns."
        )
    for row in rows_tokens:
        row[0], row[label_column] = row[label_column], row[0]


def read_row_data(
    filepath: str,
    delimiter: Optional[str] = None,
    decimal_separator: Optional[str] = None,
    header: Optional[bool] = None,
    analyze_rows: int = 20,
    zero_padding: int = 4,
    header_threshold: float = 0.5,
    label_column: Optional[int] = None,
    sheet_name: Optional[str] = None,
    header_row: Optional[int] = None,
    index_x: bool = False,
) -> List[Dict[str, Union[np.ndarray, str, Dict]]]:
    """
    Read row-oriented tabular data and return a list of spectrum dicts,
    each with keys: x_scale, y_scale, label, metadata.

    Parameters mirror `table_data_converter.read_table_data`, minus
    `interlaced_format` — row-oriented + interlaced is not supported.

    header_threshold : fraction (0.0-1.0) of non-numeric tokens required to
                        auto-detect the label column. Only used when
                        `header` is None. Default 0.5 (50%).
    label_column      : zero-based index of the raw file column to use as
                         the spectrum-label source. None → column 0
                         (the default, row-oriented-standard position).
    sheet_name        : Excel files only. Which worksheet to read.
                         None → prefer "spectra", else the first
                         sheet. Falls back to that same rule if the named
                         sheet doesn't exist in this particular file.
    """
    logger.info("Importing (row-oriented): %s", os.path.basename(filepath))

    file_ext = os.path.splitext(filepath)[1].lower()
    if file_ext in ('.xlsx', '.xls', '.xlsm'):
        return _read_row_excel_data(
            filepath, header=header, zero_padding=zero_padding,
            header_threshold=header_threshold, label_column=label_column,
            sheet_name=sheet_name, header_row=header_row, index_x=index_x,
        )

    # Local imports to avoid a circular import at module load time, and to
    # reuse the exact same, already-hardened helpers the column parser uses.
    from src.modules.data_io.table_data_converter import (
        _read_raw_data, _detect_delimiter, _detect_decimal_separator,
        _detect_header, _split_line, _parse_column_data,
    )

    raw_data = _read_raw_data(filepath)
    if not raw_data:
        raise ValueError("Empty file or no valid data found")

    # Discard any preamble above the chosen starting row (no-op when Auto).
    # In THIS layout row 0 is the shared X-scale row, so "header row" here means
    # "the row the table really starts at" — i.e. the X-scale row. The header
    # FLAG in row-oriented refers to the label COLUMN (see label_column), not to
    # a row, so it is deliberately left untouched.
    raw_data = _apply_header_row_rows(raw_data, header_row)

    logger.debug("Read %d lines", len(raw_data))

    sample = raw_data[:analyze_rows]

    # --- delimiter / decimal detection --------------------------------------
    # These operate on individual raw lines exactly as they do for column
    # data — orientation doesn't change how a single line is delimited.
    if delimiter is None:
        delimiter = _detect_delimiter(sample)
    if decimal_separator is None:
        decimal_separator = _detect_decimal_separator(sample, delimiter)

    logger.debug("Using: decimal=%r, delimiter=%r", decimal_separator, delimiter)

    # --- tokenize & transpose -------------------------------------------
    rows_tokens = [_split_line(line, delimiter) for line in raw_data if line.strip()]
    if not rows_tokens:
        raise ValueError("No data rows found")

    max_len = max(len(r) for r in rows_tokens)
    for r in rows_tokens:
        r.extend([''] * (max_len - len(r)))

    _apply_label_column(rows_tokens, label_column)

    transposed = [list(col) for col in zip(*rows_tokens)]

    # Tokens are already individual values (delimiter-split), so any
    # delimiter is safe to rejoin with — tab avoids ambiguity with any
    # separator character that might appear inside a label.
    join_delim = '\t'
    transposed_lines = [join_delim.join(tok) for tok in transposed]

    # --- header (= chosen label column) -----------------------------------
    # After transposing, the label column becomes the first line, so the
    # existing row-0 header heuristic applies unchanged.
    if header is None:
        has_header = _detect_header(transposed_lines, join_delim, decimal_separator, header_threshold)
    else:
        has_header = header
    logger.debug("Header column: %s (user specified: %s)", has_header, header is not None)

    # Synthetic 1..N X axis. In THIS layout the X-scale normally comes from the
    # first data ROW; index_x means there is no such row — every row is a
    # spectrum, and the X axis becomes the point number (1, 2, 3, …). After the
    # transpose that is exactly "prepend a synthetic first column", so the same
    # helper the standard layout uses applies unchanged.
    if index_x:
        from src.modules.data_io.table_data_converter import _prepend_index_column
        transposed_lines = _prepend_index_column(
            transposed_lines, join_delim, has_header)

    # --- parse using the existing column-format logic -----------------------
    spectra = _parse_column_data(
        transposed_lines, filepath, join_delim, decimal_separator,
        has_header, zero_padding,
    )

    import_params = {
        'delimiter': delimiter,
        'decimal_separator': decimal_separator,
        'has_header': has_header,
        'row_oriented': True,
        'label_column': label_column if label_column is not None else 0,
        'detection_mode': 'manual' if (delimiter != 'auto') else 'auto',
    }
    for spectrum in spectra:
        spectrum['metadata']['import_parameters'] = import_params
        spectrum['metadata']['file_type'] = 'row_table'

    logger.info(
        "Created %d spectra (row-oriented) from %s",
        len(spectra), os.path.basename(filepath)
    )
    return spectra


# ---------------------------------------------------------------------------
# Excel reader
# ---------------------------------------------------------------------------

def _read_row_excel_data(
    filepath: str,
    header: Optional[bool] = None,
    zero_padding: int = 4,
    sheet_name: Optional[str] = None,
    header_threshold: float = 0.5,
    label_column: Optional[int] = None,
    header_row: Optional[int] = None,
    index_x: bool = False,
) -> List[Dict[str, Union[np.ndarray, str, Dict]]]:
    """
    Read row-oriented spectra from an Excel (.xlsx / .xls / .xlsm) file.

    Same sheet-selection rules as the column-format Excel reader: prefers
    a sheet named "spectra", otherwise uses the first sheet.
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "openpyxl is required to import Excel files.\n"
            "Install it with:  pip install openpyxl"
        )

    from src.modules.data_io.table_data_converter import (
        _detect_header, _parse_column_data,
    )

    logger.info("Importing Excel (row-oriented): %s", os.path.basename(filepath))

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)

    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    elif 'spectra' in wb.sheetnames:
        ws = wb['spectra']
        logger.debug("Using sheet: spectra")
    else:
        ws = wb.active
        logger.debug("Using sheet: %s", ws.title)
    used_sheet_title = ws.title

    raw_rows: List[List[str]] = []
    for row in ws.iter_rows(values_only=True):
        tokens = []
        for cell in row:
            if cell is None:
                tokens.append('')
            elif isinstance(cell, float):
                tokens.append(repr(cell))
            elif isinstance(cell, int):
                tokens.append(str(cell))
            else:
                tokens.append(str(cell).strip())
        while tokens and tokens[-1] == '':
            tokens.pop()
        if tokens:
            raw_rows.append(tokens)

    wb.close()

    if not raw_rows:
        raise ValueError("Excel file contains no data")

    raw_rows = _apply_header_row_rows(raw_rows, header_row)

    logger.debug("Read %d non-empty rows from Excel", len(raw_rows))

    max_len = max(len(r) for r in raw_rows)
    for r in raw_rows:
        r.extend([''] * (max_len - len(r)))

    _apply_label_column(raw_rows, label_column)

    transposed = [list(col) for col in zip(*raw_rows)]
    delimiter = '\t'
    decimal_separator = '.'   # cells were converted via repr() → always dot
    transposed_lines = [delimiter.join(tok) for tok in transposed]

    if header is None:
        has_header = _detect_header(transposed_lines, delimiter, decimal_separator, header_threshold)
    else:
        has_header = header
    logger.debug("Header column: %s", has_header)

    # Synthetic 1..N X axis — see the note in read_row_data().
    if index_x:
        from src.modules.data_io.table_data_converter import _prepend_index_column
        transposed_lines = _prepend_index_column(
            transposed_lines, delimiter, has_header)

    spectra = _parse_column_data(
        transposed_lines, filepath, delimiter, decimal_separator,
        has_header, zero_padding,
    )

    import_params = {
        'delimiter': 'tab (Excel)',
        'decimal_separator': '.',
        'has_header': has_header,
        'row_oriented': True,
        'label_column': label_column if label_column is not None else 0,
        'sheet_name': used_sheet_title,
        'detection_mode': 'excel',
    }
    for spectrum in spectra:
        spectrum['metadata']['import_parameters'] = import_params
        spectrum['metadata']['file_type'] = 'row_table_excel'

    logger.info(
        "Created %d spectra (row-oriented Excel) from %s",
        len(spectra), os.path.basename(filepath)
    )
    return spectra
