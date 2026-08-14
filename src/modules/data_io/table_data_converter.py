# src/modules/data_io/table_data_converter.py
#
# Column-oriented import for spectra files.
#
# Expected format (standard):
#   col 0      col 1 … col N
#   x_value    y1      y2  …
#   …
#
# Interlaced format:
#   x1  y1  x2  y2  …
#
# Optional first row header (auto-detected or forced via `header` parameter).
#
# Excel (.xlsx) support:
#   Sheet cells are converted to tab-separated strings and fed into the
#   same parsing pipeline as plain-text files.  The active sheet (index 0)
#   is used unless the file contains a sheet named "spectra" (the name
#   written by SaveManager), which is preferred.

import numpy as np
from typing import List, Dict, Union, Optional, Tuple
from datetime import datetime
import os
import re

from src.modules.utils.app_logger import get_logger
logger = get_logger(__name__)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------



def _prepend_index_column(raw_data: list, delimiter: str, has_header: bool) -> list:
    """
    Insert a synthetic X column holding 1, 2, 3, … in front of the data.

    Some tables have no meaningful numeric X axis at all — concentration
    profiles, for instance, where the only non-numeric column is a sample name.
    Excluding that name column doesn't help: the FIRST remaining column would
    then be eaten as the X-scale, silently costing you a real data column.

    Injecting a synthetic X column instead means every real column stays a
    spectrum, and the X axis becomes the row number. It is done by rewriting the
    raw lines rather than by special-casing the parser, so everything downstream
    (header handling, labels, NaN skipping) keeps working exactly as before.
    """
    out = []
    counter = 0
    for i, line in enumerate(raw_data):
        if i == 0 and has_header:
            out.append('index' + delimiter + line)
            continue
        counter += 1
        out.append(str(counter) + delimiter + line)
    return out


def _apply_header_row(rows: list, header_row: Optional[int]):
    """
    Drop everything ABOVE the user-chosen header row, so that row becomes row 0.

    Files often carry preamble above the real table — instrument banners, blank
    lines, a title, units on their own line. Previously the header was always
    assumed to be the FIRST row, so such files could not be imported correctly
    at all.

    One mechanism serves both layouts, because "the row the table really starts
    at" means the same thing in each:

      * Standard (spectra in columns) — that row holds the column names, so
        after slicing, the existing "header is row 0" logic applies unchanged,
        and header is forced True (the user has told us it IS a header).
      * Row-oriented (spectra in rows) — row 0 is the shared X-scale row, so
        slicing makes the chosen row the X-scale row. The header flag there
        refers to the label COLUMN (see label_column), not to a row, so it is
        deliberately left alone.

    Returns the sliced rows.
    """
    if header_row is None:
        return rows
    if header_row < 0 or header_row >= len(rows):
        raise ValueError(
            f"Header row {header_row + 1} is out of range — the file only has "
            f"{len(rows)} rows."
        )
    return rows[header_row:]


def read_table_data(
    filepath: str,
    delimiter: Optional[str] = None,
    decimal_separator: Optional[str] = None,
    header: Optional[bool] = None,
    analyze_rows: int = 20,
    zero_padding: int = 4,
    interlaced_format: bool = False,
    row_oriented: bool = False,
    header_threshold: float = 0.5,
    label_column: Optional[int] = None,
    x_scale_column: Optional[int] = None,
    exclude_columns: Optional[List[int]] = None,
    sheet_name: Optional[str] = None,
    header_row: Optional[int] = None,
    index_x: bool = False,
) -> List[Dict[str, Union[np.ndarray, str, Dict]]]:
    """
    Read column-oriented, interlaced, or row-oriented tabular data and
    return a list of spectrum dicts, each with keys:
    x_scale, y_scale, label, metadata.

    Parameters
    ----------
    filepath          : path to the input file
    delimiter         : column separator; None → auto-detect
                        (ignored for .xlsx — cells are always tab-joined)
    decimal_separator : '.' or ','; None → auto-detect
                        (ignored for .xlsx — numbers are read as floats)
    header            : True / False / None (None → auto-detect)
    analyze_rows      : how many rows to use for auto-detection heuristics
    zero_padding      : digit width for auto-generated spectrum numbers
    interlaced_format : if True, columns are x1,y1,x2,y2,… pairs
    row_oriented      : if True, spectra are laid out in rows instead of
                         columns (row 0 = shared X-scale, an optional first
                         column provides spectrum labels). Mutually
                         exclusive with interlaced_format.
    header_threshold  : fraction (0.0-1.0) of non-numeric tokens in the
                         candidate header line/column required to auto-detect
                         it as a header. Only used when `header` is None.
                         Default 0.5 (i.e. 50%).
    label_column      : row-oriented only. Zero-based index of the column
                         (in the raw file, before any transposition) to use
                         as the spectrum-label source, instead of the
                         default column 0. None → column 0.
    x_scale_column    : Standard layout only. Zero-based index (in the raw
                         file) of the column to use as the shared X-scale,
                         instead of the default column 0. None → the first
                         column remaining after `exclude_columns` is
                         removed. Mutually exclusive with interlaced_format
                         and row_oriented. Cannot be one of exclude_columns.
    exclude_columns   : Standard layout only. Zero-based indices (in the
                         raw file) of columns to drop entirely before
                         parsing — for metadata columns (IDs, sample types,
                         timestamps, …) that are neither the X-scale nor a
                         real spectrum. Unlike swapping the X-scale column
                         alone, excluded columns never become bogus
                         spectra. None/empty → nothing excluded.
    sheet_name        : Excel files only (.xlsx/.xls/.xlsm), ignored for
                         plain text. Which worksheet to read. None → prefer
                         a sheet named "spectra" (the sheet this
                         app's own Save function writes), else the first
                         sheet. If the named sheet doesn't exist in this
                         particular file, silently falls back to the same
                         Auto rule rather than raising.
    """

    logger.info("Importing: %s", os.path.basename(filepath))

    if row_oriented and interlaced_format:
        raise ValueError(
            "row_oriented and interlaced_format cannot both be enabled — "
            "choose one layout."
        )

    if index_x and x_scale_column is not None:
        raise ValueError(
            "index_x and x_scale_column cannot both be set — index_x replaces "
            "the X axis with the row number, so there is no X column to pick."
        )

    if x_scale_column is not None and (interlaced_format or row_oriented):
        raise ValueError(
            "x_scale_column is only supported in Standard layout — it "
            "cannot be combined with interlaced_format or row_oriented."
        )

    if exclude_columns and (interlaced_format or row_oriented):
        raise ValueError(
            "exclude_columns is only supported in Standard layout — it "
            "cannot be combined with interlaced_format or row_oriented."
        )

    if exclude_columns and x_scale_column is not None and x_scale_column in exclude_columns:
        raise ValueError(
            f"x_scale_column={x_scale_column} cannot also be in "
            f"exclude_columns — pick a different X-scale column, or stop "
            f"excluding this one."
        )

    if row_oriented:
        from src.modules.data_io.row_data_converter import read_row_data
        return read_row_data(
            filepath,
            delimiter=delimiter,
            decimal_separator=decimal_separator,
            header=header,
            analyze_rows=analyze_rows,
            zero_padding=zero_padding,
            header_threshold=header_threshold,
            label_column=label_column,
            sheet_name=sheet_name,
            header_row=header_row,
            index_x=index_x,
        )

    file_ext = os.path.splitext(filepath)[1].lower()

    # --- Excel path ---------------------------------------------------------
    if file_ext in ('.xlsx', '.xls', '.xlsm'):
        return read_excel_data(
            filepath,
            header=header,
            zero_padding=zero_padding,
            interlaced_format=interlaced_format,
            header_threshold=header_threshold,
            x_scale_column=x_scale_column,
            exclude_columns=exclude_columns,
            sheet_name=sheet_name,
            header_row=header_row,
            index_x=index_x,
        )

    # --- Plain-text path ----------------------------------------------------
    raw_data = _read_raw_data(filepath)
    if not raw_data:
        raise ValueError("Empty file or no valid data found")

    # Discard any preamble above the user-chosen header row (no-op when Auto).
    raw_data = _apply_header_row(raw_data, header_row)
    if header_row is not None:
        # The user has explicitly pointed at the header row, so it IS a header.
        header = True

    logger.debug("Read %d lines", len(raw_data))

    sample = raw_data[:analyze_rows]

    # --- delimiter detection ------------------------------------------------
    if delimiter is None:
        delimiter = _detect_delimiter(sample)

    # --- decimal separator detection ----------------------------------------
    if decimal_separator is None:
        decimal_separator = _detect_decimal_separator(sample, delimiter)

    logger.debug("Using: decimal=%r, delimiter=%r", decimal_separator, delimiter)

    # --- column selection (exclude + X-scale) --------------------------------
    # Drop excluded columns, then swap the chosen X-scale column into
    # position 0, so the existing column-0-is-X-scale logic (header
    # detection, _parse_column_data) applies unchanged — same technique
    # used for label_column in row-oriented imports.
    effective_x_scale_column = x_scale_column
    if not interlaced_format and (exclude_columns or x_scale_column not in (None, 0)):
        raw_data, delimiter, effective_x_scale_column = _apply_standard_column_selection(
            raw_data, delimiter, x_scale_column, exclude_columns
        )

    # --- header detection ---------------------------------------------------
    if header is None:
        has_header = _detect_header(raw_data, delimiter, decimal_separator, header_threshold)
    else:
        has_header = header
    logger.debug("Header: %s (user specified: %s)", has_header, header is not None)

    # Synthetic 1..N X axis, injected only AFTER the header row is known so the
    # header gets a name and the data rows get numbers.
    if index_x and not interlaced_format:
        raw_data = _prepend_index_column(raw_data, delimiter, has_header)

    # --- parse --------------------------------------------------------------
    if interlaced_format:
        from src.modules.data_io.interlaced_data_converter import read_interlaced_data
        spectra = read_interlaced_data(
            filepath, delimiter, decimal_separator,
            has_header, analyze_rows, zero_padding,
            header_threshold=header_threshold,
        )
    else:
        spectra = _parse_column_data(
            raw_data, filepath, delimiter, decimal_separator, has_header, zero_padding
        )

    # Stamp import parameters on every spectrum
    import_params = {
        'delimiter': delimiter,
        'decimal_separator': decimal_separator,
        'has_header': has_header,
        'interlaced_format': interlaced_format,
        'x_scale_column': effective_x_scale_column if not interlaced_format else None,
        'exclude_columns': sorted(exclude_columns) if exclude_columns else [],
        'detection_mode': 'manual' if (delimiter != 'auto') else 'auto',
    }
    for spectrum in spectra:
        spectrum['metadata']['import_parameters'] = import_params

    logger.info("Created %d spectra from %s", len(spectra), os.path.basename(filepath))
    return spectra


# ---------------------------------------------------------------------------
# Excel reader
# ---------------------------------------------------------------------------

def read_excel_data(
    filepath: str,
    header: Optional[bool] = None,
    zero_padding: int = 4,
    interlaced_format: bool = False,
    sheet_name: Optional[str] = None,
    header_threshold: float = 0.5,
    x_scale_column: Optional[int] = None,
    exclude_columns: Optional[List[int]] = None,
    header_row: Optional[int] = None,
    index_x: bool = False,
) -> List[Dict[str, Union[np.ndarray, str, Dict]]]:
    """
    Read spectra from an Excel (.xlsx / .xls / .xlsm) file.

    The preferred sheet is "spectra" (written by SaveManager).
    If that sheet is absent, the first sheet is used.

    Numbers are read as native floats so no decimal-separator ambiguity
    exists.  The resulting data is converted to tab-separated text rows
    and fed into the standard column / interlaced parser so all existing
    logic is reused.

    Parameters
    ----------
    filepath        : path to the Excel file
    header          : True / False / None (auto-detect)
    zero_padding    : digit width for auto-generated spectrum numbers
    interlaced_format : route to interlaced parser when True
    sheet_name      : explicit sheet name; None → auto-select. If the
                      given name doesn't exist in this file, silently
                      falls back to the same Auto rule rather than
                      raising — useful when applying one sheet_name
                      across a batch of files that don't all share it,
                      at the cost of not erroring on a genuine typo.
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "openpyxl is required to import Excel files.\n"
            "Install it with:  pip install openpyxl"
        )

    logger.info("Importing Excel: %s", os.path.basename(filepath))

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)

    # Select sheet
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    elif 'spectra' in wb.sheetnames:
        ws = wb['spectra']
        logger.debug("Using sheet: spectra")
    else:
        ws = wb.active
        logger.debug("Using sheet: %s", ws.title)
    used_sheet_title = ws.title

    # Read all rows, converting each cell to a string token.
    # Empty trailing cells are stripped; completely empty rows are skipped.
    raw_rows: List[List[str]] = []
    for row in ws.iter_rows(values_only=True):
        tokens = []
        for cell in row:
            if cell is None:
                tokens.append('')
            elif isinstance(cell, float):
                # Represent floats precisely without locale issues
                tokens.append(repr(cell))
            elif isinstance(cell, int):
                tokens.append(str(cell))
            else:
                tokens.append(str(cell).strip())
        # Strip trailing empty tokens
        while tokens and tokens[-1] == '':
            tokens.pop()
        if tokens:
            raw_rows.append(tokens)

    wb.close()

    if not raw_rows:
        raise ValueError("Excel file contains no data")

    logger.debug("Read %d non-empty rows from Excel", len(raw_rows))

    # Column selection (exclude + X-scale), Standard layout only — same
    # combined logic as the plain-text path, operating directly on the
    # already-tokenized rows (no need to re-split by delimiter here).
    effective_x_scale_column = x_scale_column
    if not interlaced_format and (exclude_columns or x_scale_column not in (None, 0)):
        max_len = max(len(r) for r in raw_rows)
        for r in raw_rows:
            r.extend([''] * (max_len - len(r)))

        exclude_set = set(exclude_columns) if exclude_columns else set()
        for idx in exclude_set:
            if idx >= max_len:
                raise ValueError(
                    f"exclude_columns contains {idx}, which is out of "
                    f"range — the sheet has only {max_len} columns."
                )
        if x_scale_column is not None:
            if x_scale_column >= max_len:
                raise ValueError(
                    f"x_scale_column={x_scale_column} is out of range — "
                    f"the sheet has only {max_len} columns."
                )
            if x_scale_column in exclude_set:
                raise ValueError(
                    f"x_scale_column={x_scale_column} cannot also be in "
                    f"exclude_columns."
                )

        kept_indices = [i for i in range(max_len) if i not in exclude_set]
        if not kept_indices:
            raise ValueError("exclude_columns removes every column in the sheet.")

        effective_x_scale_column = x_scale_column if x_scale_column is not None else kept_indices[0]
        new_x_position = kept_indices.index(effective_x_scale_column)

        raw_rows = [[r[i] for i in kept_indices] for r in raw_rows]
        for r in raw_rows:
            r[0], r[new_x_position] = r[new_x_position], r[0]

    # Convert to tab-separated strings — the existing parsers all accept this
    raw_data = ['\t'.join(row) for row in raw_rows]

    # Drop any preamble above the chosen header row (no-op when Auto), exactly
    # as for plain text — Excel sheets carry banners/titles above the table too.
    raw_data = _apply_header_row(raw_data, header_row)
    if header_row is not None:
        header = True
    delimiter = '\t'
    decimal_separator = '.'   # cells were converted via repr() → always dot

    # Header detection
    if header is None:
        has_header = _detect_header(raw_data, delimiter, decimal_separator, header_threshold)
    else:
        has_header = header

    # Synthetic 1..N X axis (see _prepend_index_column) — injected after the
    # header row is known, same as the plain-text path.
    if index_x:
        raw_data = _prepend_index_column(raw_data, delimiter, has_header)
    logger.debug("Header: %s", has_header)

    # Parse using existing logic
    if interlaced_format:
        # For interlaced Excel we re-use the interlaced converter but pass
        # the already-detected delimiter / decimal so no re-detection occurs.
        from src.modules.data_io.interlaced_data_converter import read_interlaced_data
        spectra = _parse_interlaced_from_raw(
            raw_data, filepath, delimiter, decimal_separator,
            has_header, zero_padding
        )
    else:
        spectra = _parse_column_data(
            raw_data, filepath, delimiter, decimal_separator, has_header, zero_padding
        )

    # Stamp import parameters
    import_params = {
        'delimiter': 'tab (Excel)',
        'decimal_separator': '.',
        'has_header': has_header,
        'interlaced_format': interlaced_format,
        'x_scale_column': effective_x_scale_column if not interlaced_format else None,
        'exclude_columns': sorted(exclude_columns) if exclude_columns else [],
        'sheet_name': used_sheet_title,
        'detection_mode': 'excel',
    }
    for spectrum in spectra:
        spectrum['metadata']['import_parameters'] = import_params
        spectrum['metadata']['file_type'] = 'excel'

    logger.info("Created %d spectra from Excel: %s", len(spectra), os.path.basename(filepath))
    return spectra


def _parse_interlaced_from_raw(
    raw_data: List[str],
    filepath: str,
    delimiter: str,
    decimal_separator: str,
    has_header: bool,
    zero_padding: int,
) -> List[Dict]:
    """
    Thin wrapper: pass already-converted raw_data directly to the
    interlaced parser without triggering file re-read.
    """
    # The interlaced converter re-reads the file from disk; to avoid that
    # we reuse _parse_column_data logic but with paired columns.
    # This mirrors _parse_interlaced_column_data from interlaced_data_converter
    # but operates on our already-loaded raw_data list.
    labels: List[str] = []
    data_lines = raw_data.copy()

    if has_header and data_lines:
        labels = _split_line(data_lines[0], delimiter)
        data_lines = data_lines[1:]

    # Build numeric array
    all_rows: List[List[float]] = []
    for line in data_lines:
        if not line.strip():
            continue
        parts = _split_line(line, delimiter)
        numbers = [_safe_convert(p, decimal_separator) for p in parts]
        if len(numbers) >= 2 and not np.isnan(numbers[0]):
            all_rows.append(numbers)

    if not all_rows:
        raise ValueError("No valid numeric data found in Excel file")

    max_cols = max(len(r) for r in all_rows)
    for row in all_rows:
        row.extend([np.nan] * (max_cols - len(row)))

    data_array = np.array(all_rows, dtype=float)
    base_name = os.path.splitext(os.path.basename(filepath))[0]
    num_pairs = data_array.shape[1] // 2
    spectra: List[Dict] = []

    # Cache file timestamps ONCE, same as _parse_column_data — this
    # function previously referenced file_ctime/file_mtime without ever
    # defining them, which raised NameError on every Excel interlaced
    # import that produced at least one valid spectrum (i.e. always).
    try:
        file_ctime = datetime.fromtimestamp(os.path.getctime(filepath)).isoformat()
        file_mtime = datetime.fromtimestamp(os.path.getmtime(filepath)).isoformat()
    except OSError:
        file_ctime = ''
        file_mtime = ''

    for pair_idx in range(num_pairs):
        x_col = pair_idx * 2
        y_col = pair_idx * 2 + 1
        if y_col >= data_array.shape[1]:
            break

        x_values = data_array[:, x_col]
        y_values = data_array[:, y_col]
        valid = ~(np.isnan(x_values) | np.isnan(y_values))
        x_f = x_values[valid]
        y_f = y_values[valid]

        if len(x_f) < 2:
            continue

        # Label
        if labels and y_col < len(labels) and labels[y_col].strip():
            clean = re.sub(r'[^\w\s\-\.]', '_', labels[y_col].strip())
            clean = re.sub(r'\s+', '_', clean)
            label = f"{base_name} : {clean}"
        elif num_pairs == 1:
            label = base_name
        else:
            label = f"{base_name} : {str(pair_idx + 1).zfill(zero_padding)}"

        metadata = {
            'file_path': filepath,
            'file_type': 'excel_interlaced',
            'original_label': label,
            'spectrum_name': label,
            'interlaced_pair_index': pair_idx,
            'valid_points': len(x_f),
            'file_creation_time': file_ctime,
            'file_modification_time': file_mtime,
        }

        spectra.append({
            'x_scale': x_f.astype(float),
            'y_scale': y_f.astype(float),
            'label': label,
            'metadata': metadata,
        })
        logger.debug("Created interlaced spectrum %r: %d points", label, len(x_f))

    if not spectra:
        raise ValueError("No valid spectra created from Excel interlaced data")
    return spectra


# ---------------------------------------------------------------------------
# File reading
# ---------------------------------------------------------------------------

def _apply_standard_column_selection(
    raw_data: List[str],
    delimiter: str,
    x_scale_column: Optional[int],
    exclude_columns: Optional[List[int]],
) -> Tuple[List[str], str, int]:
    """
    Drop `exclude_columns` entirely, then swap the chosen X-scale column
    into position 0 across every line — so the rest of the pipeline
    (header detection, _parse_column_data, which both hardcode "column 0
    is X-scale") applies unchanged, the same technique `row_data_converter`
    uses for `label_column`.

    Both `x_scale_column` and `exclude_columns` refer to ORIGINAL column
    indices, i.e. positions in the raw file before anything is removed —
    matching what the dialog's column picker shows the user. Excluded
    columns are removed first; if `x_scale_column` is None, "Auto" means
    "the first column remaining after exclusion" (mirroring how Origin/
    Igor-style "Disregard" columns simply vanish and roles apply to
    whatever is left), not literally raw index 0 if that happens to have
    been excluded.

    Lines are rejoined with tab, and tab is returned as the new delimiter,
    since tokens are already individual values at this point and tab
    avoids ambiguity with whatever character the original delimiter was.

    Returns (new_lines, new_delimiter, effective_x_scale_column) where
    effective_x_scale_column is the ORIGINAL index actually used as X —
    useful for stamping accurate import metadata.

    Raises ValueError if any index is out of range, or if x_scale_column
    is itself one of exclude_columns.
    """
    rows_tokens = [_split_line(line, delimiter) for line in raw_data]
    if not rows_tokens:
        return raw_data, delimiter, (x_scale_column if x_scale_column is not None else 0)

    max_len = max(len(r) for r in rows_tokens)
    for r in rows_tokens:
        r.extend([''] * (max_len - len(r)))

    exclude_set = set(exclude_columns) if exclude_columns else set()
    for idx in exclude_set:
        if idx >= max_len:
            raise ValueError(
                f"exclude_columns contains {idx}, which is out of range — "
                f"the file has only {max_len} columns."
            )

    if x_scale_column is not None:
        if x_scale_column >= max_len:
            raise ValueError(
                f"x_scale_column={x_scale_column} is out of range — the "
                f"file has only {max_len} columns."
            )
        if x_scale_column in exclude_set:
            raise ValueError(
                f"x_scale_column={x_scale_column} cannot also be in "
                f"exclude_columns."
            )

    kept_indices = [i for i in range(max_len) if i not in exclude_set]
    if not kept_indices:
        raise ValueError("exclude_columns removes every column in the file.")

    effective_x_scale_column = x_scale_column if x_scale_column is not None else kept_indices[0]
    new_x_position = kept_indices.index(effective_x_scale_column)

    filtered_rows = [[r[i] for i in kept_indices] for r in rows_tokens]
    for r in filtered_rows:
        r[0], r[new_x_position] = r[new_x_position], r[0]

    join_delim = '\t'
    new_lines = [join_delim.join(r) for r in filtered_rows]
    return new_lines, join_delim, effective_x_scale_column


def _read_raw_data(filepath: str) -> List[str]:
    """
    Read file lines, trying common encodings; strips blank lines and BOM.

    IMPORTANT: uses rstrip('\\r\\n') (not strip()) so that leading
    whitespace/tabs are preserved. A leading tab is meaningful — it's an
    empty first cell, e.g. a blank corner cell above the X-scale column
    in a header row (a very common convention), or a leading empty field
    at the start of a row. Calling strip() would silently eat that
    leading tab and shift every column index left by one for that line,
    misaligning header labels with the wrong Y-column for the rest of the
    file — a real, silent mislabeling bug this fixes. (Same reasoning and
    fix as interlaced_data_converter._read_raw_data.)
    """
    for encoding in ('utf-8-sig', 'utf-8', 'latin-1', 'cp1252'):
        try:
            lines = []
            with open(filepath, 'r', encoding=encoding) as fh:
                for line in fh:
                    line = line.rstrip('\r\n')
                    if line.startswith('\ufeff'):
                        line = line[1:]
                    if line.strip():
                        lines.append(line)
            if lines:
                return lines
        except UnicodeDecodeError:
            continue

    raise ValueError("Could not read file with any supported encoding")


# ---------------------------------------------------------------------------
# Separator detection  (Bug 5 fix: delimiter detected before decimal,
#  and comma is only chosen as delimiter when no decimal-comma evidence exists)
# ---------------------------------------------------------------------------

def _detect_delimiter(lines: List[str]) -> str:
    """
    Detect the column delimiter.

    Priority: tab → semicolon → pipe → comma (only if comma cannot be
    the decimal separator) → whitespace.

    Tab detection uses the number of *columns* (split result length) rather
    than raw character count, because tab-delimited files with long numeric
    values still have a consistent column count even when character counts
    vary slightly.
    """
    if not lines:
        return ','

    non_empty = [line for line in lines if line.strip()]

    # --- Tab: check column count consistency, not character count ------------
    # Also require that a real majority of lines actually contain a tab —
    # otherwise a single stray tab character (e.g. left over from a
    # copy/paste, or trailing whitespace before a line ending — confirmed
    # to happen in real files) makes max(tab_col_counts) > 1 for just that
    # one line while every other line has none, and the old
    # len(set(...)) <= 2 tolerance (meant for a ragged trailing column)
    # would still accept it as "tab-delimited", causing every genuinely
    # space-delimited line to fail to split at all during parsing.
    tab_col_counts = [len(line.split('\t')) for line in non_empty]
    lines_with_tab = sum(1 for c in tab_col_counts if c > 1)
    if (tab_col_counts and max(tab_col_counts) > 1 and len(set(tab_col_counts)) <= 2
            and lines_with_tab >= max(1, len(non_empty) * 0.5)):
        logger.debug("Selected delimiter: tab")
        return '\t'

    # --- Semicolon and pipe: character count consistency --------------------
    def _consistent(delim: str) -> bool:
        counts = [line.count(delim) for line in non_empty]
        if not counts or max(counts) == 0:
            return False
        return len(set(counts)) <= 2

    for delim in (';', '|'):
        if _consistent(delim):
            logger.debug("Selected delimiter: %r", delim)
            return delim

    # Comma: only treat as delimiter if it does NOT look like a decimal separator.
    # Heuristic: if every number token that contains a comma has exactly one
    # comma flanked by digits on both sides → it is a decimal separator, not a
    # column delimiter.
    if lines[0].count(',') > 0:
        # Count tokens that look like "digit,digit" (decimal use)
        decimal_comma = sum(
            len(re.findall(r'\d,\d', line)) for line in lines
        )
        # Count commas that separate whole tokens (delimiter use):
        # a delimiter comma is typically preceded/followed by whitespace or
        # is at the start/end of a field.
        delimiter_comma = sum(line.count(',') for line in lines) - decimal_comma
        if _consistent(',') and delimiter_comma > decimal_comma:
            logger.debug("Selected delimiter: comma")
            return ','

    # Whitespace fallback
    first_parts = len(lines[0].split()) if lines else 0
    if first_parts > 1:
        if all(
            len(line.split()) == first_parts
            for line in lines[:min(5, len(lines))]
            if line.strip()
        ):
            logger.debug("Selected delimiter: whitespace")
            return r'\s+'

    logger.debug("Fallback delimiter: comma")
    return ','


def _detect_decimal_separator(lines: List[str], delimiter: str) -> str:
    """Count dot-decimals vs comma-decimals in numeric tokens."""
    dot_count = comma_count = 0

    for line in lines[:min(10, len(lines))]:
        parts = _split_line(line, delimiter)
        for part in parts[:10]:
            part = part.strip()
            if re.search(r'\d\.\d', part):
                dot_count += 1
            if re.search(r'\d,\d', part):
                comma_count += 1

    logger.debug("Decimal counts: dots=%d, commas=%d", dot_count, comma_count)
    result = ',' if comma_count > dot_count else '.'
    logger.debug("Selected decimal separator: %r", result)
    return result


# ---------------------------------------------------------------------------
# Header detection
# ---------------------------------------------------------------------------

def _detect_header(lines: List[str], delimiter: str, decimal_separator: str,
                    threshold: float = 0.5) -> bool:
    """Return True when the first line looks like a text header.

    threshold : fraction (0.0-1.0) of tokens that must be non-numeric for
                the line to be classified as a header. Default 0.5 (50%).
    """
    if not lines:
        return False

    parts = _split_line(lines[0], delimiter)
    non_numeric = 0
    for part in parts:
        part = part.strip()
        if not part:
            continue
        try:
            _convert_to_float(part, decimal_separator)
        except ValueError:
            non_numeric += 1

    has_header = len(parts) > 0 and non_numeric >= len(parts) * threshold
    logger.debug("Header check: %d/%d non-numeric parts (threshold=%.0f%%) → %s",
                 non_numeric, len(parts), threshold * 100, has_header)
    return has_header


# ---------------------------------------------------------------------------
# Low-level utilities
# ---------------------------------------------------------------------------

def _split_line(line: str, delimiter: str) -> List[str]:
    """Split a line on the given delimiter."""
    if delimiter == r'\s+':
        return [t.strip() for t in line.split() if t.strip()]
    return [t.strip() for t in line.split(delimiter)]


def _convert_to_float(value: str, decimal_separator: str) -> float:
    """
    Convert a string token to float.

    Handles:
    - comma or dot as decimal separator
    - scientific notation: 1.5e3, 1.5E+3, 1,5e-3, etc.

    Raises ValueError for non-numeric strings (so callers can distinguish
    NaN-by-content from parse failure).
    """
    clean = value.strip()
    if not clean:
        return np.nan

    if decimal_separator == ',':
        # Replace decimal comma with dot, but do NOT touch the exponent sign
        # e.g. "1,5E+3" → "1.5E+3"
        clean = re.sub(r'(\d),(\d)', r'\1.\2', clean)

    # Normalise scientific notation edge cases Python's float() already handles
    # most, but "1.5e3" (no sign) is valid and works; we just need to not
    # accidentally mangle things.
    return float(clean)  # raises ValueError if unparseable


def _safe_convert(value: str, decimal_separator: str) -> float:
    """Like _convert_to_float but returns np.nan instead of raising."""
    try:
        return _convert_to_float(value, decimal_separator)
    except (ValueError, TypeError):
        return np.nan


# ---------------------------------------------------------------------------
# Column data parser  (Bug 3 fix: spectrum counter is independent of col_idx)
# ---------------------------------------------------------------------------

def _parse_column_data(
    raw_data: List[str],
    filepath: str,
    delimiter: str,
    decimal_separator: str,
    has_header: bool,
    zero_padding: int,
) -> List[Dict]:
    """
    Parse standard column-oriented data:
      column 0 → shared X-scale
      columns 1…N → one spectrum each

    If a header row is present its labels are used as spectrum names.
    Completely-NaN Y-columns are silently skipped.
    Spectrum numbering uses a dedicated counter so gaps from skipped columns
    do not appear in the labels.
    """
    data_lines = raw_data.copy()
    labels: List[str] = []

    if has_header and data_lines:
        labels = _split_line(data_lines[0], delimiter)
        data_lines = data_lines[1:]
        logger.debug("Header labels: %s", labels)

    # --- cache file timestamps ONCE outside the per-spectrum loop -----------
    try:
        file_ctime = datetime.fromtimestamp(os.path.getctime(filepath)).isoformat()
        file_mtime = datetime.fromtimestamp(os.path.getmtime(filepath)).isoformat()
    except OSError:
        file_ctime = ''
        file_mtime = ''

    # --- parse numeric rows -------------------------------------------------
    all_rows: List[List[float]] = []
    for line_num, line in enumerate(data_lines):
        if not line.strip():
            continue
        parts = _split_line(line, delimiter)
        if len(parts) < 2:
            continue

        numbers = [_safe_convert(p, decimal_separator) for p in parts]
        valid_count = sum(1 for v in numbers if not np.isnan(v))

        if valid_count >= 2 and not np.isnan(numbers[0]):
            all_rows.append(numbers)
        elif line_num < 5:
            logger.debug(
                "Skipping line %d: %d valid values, X valid=%s",
                line_num + 1, valid_count, not np.isnan(numbers[0])
            )

    if not all_rows:
        logger.error("No valid numeric data rows found — every row was skipped")
        for i, line in enumerate(data_lines[:5]):
            logger.debug("  Line %d: %s", i + 1, line[:120])
        raise ValueError(
            "No valid numeric data found — every row was skipped during "
            "parsing. This usually means the file's actual layout doesn't "
            "match the selected Layout option (e.g. row-oriented or "
            "interlaced data imported as Standard, or vice versa) — check "
            "Layout and the preview, then try again."
        )

    logger.debug("Found %d valid data rows", len(all_rows))

    # Pad to rectangular array
    max_cols = max(len(row) for row in all_rows)
    for row in all_rows:
        row.extend([np.nan] * (max_cols - len(row)))

    data_array = np.array(all_rows, dtype=float)
    logger.debug("Data shape: %s", data_array.shape)

    if data_array.shape[1] < 2:
        raise ValueError("Need at least 2 columns (X + one Y)")

    # --- shared X-scale -----------------------------------------------------
    x_col = data_array[:, 0]
    valid_x = ~np.isnan(x_col)
    if not np.any(valid_x):
        raise ValueError("No valid X values in first column")
    x_scale = x_col[valid_x]
    logger.debug("X-scale: %d valid points", len(x_scale))

    # --- build spectra -------------------------------------------------------
    base_name = os.path.splitext(os.path.basename(filepath))[0]
    n_y_cols = data_array.shape[1] - 1   # total Y columns (some may be skipped)
    spectra: List[Dict] = []
    spectrum_counter = 0  # counts only spectra actually created  ← Bug 3 fix

    for col_idx in range(1, data_array.shape[1]):
        y_col = data_array[:, col_idx][valid_x]

        if np.all(np.isnan(y_col)):
            continue  # entirely empty column — skip without incrementing counter

        valid_pairs = ~np.isnan(y_col)
        if not np.any(valid_pairs):
            continue

        final_x = x_scale[valid_pairs].astype(float)
        final_y = y_col[valid_pairs].astype(float)

        if len(final_x) < 1:
            continue

        spectrum_counter += 1

        # --- label ----------------------------------------------------------
        if labels and col_idx < len(labels) and labels[col_idx].strip():
            clean = re.sub(r'[^\w\s\-\.]', '_', labels[col_idx].strip())
            clean = re.sub(r'\s+', '_', clean)
            label = f"{base_name} : {clean}"
        elif n_y_cols == 1:
            # Single spectrum — just use the file name
            label = base_name
        else:
            # Use the running counter, not col_idx  ← Bug 3 fix
            label = f"{base_name} : {str(spectrum_counter).zfill(zero_padding)}"

        metadata = {
            'file_path': filepath,
            'file_type': 'column_table',
            'original_label': label,
            'spectrum_name': label,
            'column_index': col_idx,
            'valid_points': len(final_x),
            'file_creation_time': file_ctime,
            'file_modification_time': file_mtime,
        }

        spectra.append({
            'x_scale': final_x,
            'y_scale': final_y,
            'label': label,
            'metadata': metadata,
        })

        logger.debug("Created spectrum %r: %d points", label, len(final_x))

    if not spectra:
        raise ValueError("No valid spectra created")

    return spectra
