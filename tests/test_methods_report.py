"""Methods report for publications (File -> Methods report for publication)."""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import glob
import re

import numpy as np
import pytest
from PyQt5.QtWidgets import QApplication, QFileDialog, QMessageBox

from src.modules.misc import methods_report as mr

app = QApplication.instance() or QApplication([])
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TBA = os.path.join(ROOT, 'resources', 'test_data', 'real', 'CD_spectra_TBA.txt')

# Settings dicts as the real dialogs produce them (captured from the app).
REAL_SETTINGS = {
    'Automated Baseline': {'algorithm': 'arpls', 'lambda': 1e5, 'n_iter': 50, 'fitting_ranges': [],
                           'invert_regions': False, 'processing_mode': 'auto', 'max_workers': 0},
    'SNIP Baseline': {'n_iter': 40, 'decreasing': True, 'transform': True, 'smooth_window': 0},
    'SG-smoothing': {'window_length': 11, 'polyorder': 3, 'deriv_order': 0, 'mode': 'interp',
                     'delta_values': {}, 'delta': 1.0},
    'Normalization': {'normalization_mode': 'vector', 'x_ranges': [(1000, 1100)], 'is_exclude_mode': False,
                      'x_min': 1000, 'x_max': 1100, 'reference_spectrum_index': 0, 'percentile_value': 95.0,
                      'msc_use_region': False, 'target_mode': 'standard', 'target_value': 1.0},
    'Data range': {'ranges': [(220.0, 320.0)], 'is_exclude_mode': False, 'apply_linearization': True,
                   'x_step': 0.5, 'x_min': 220.0, 'x_max': 320.0},
    'Manual baseline': {'corrections': {'a': {'type': 'polynomial', 'order': 3, 'points': [1, 2, 3, 4]},
                                        'b': {'type': 'polynomial', 'order': 3, 'points': 5}}},
    'SVD background': {'correction_mode': 'all_subspectra', 'selected_subspectra': [0, 1, 2],
                       'baseline_corrections': {}, 'inverted_subspectra': []},
    'Spike removal': {'threshold': 7.0, 'half_window': 2, 'mode': 'interactive',
                      'effective_spikes': {'a': [3]}, 'half_window_map': {'a': 2}},
    'Cosmic ray removal': {'threshold': 10.0, 'replace_window': 1},
    'FFT Denoising': {'low_cutoff': None, 'high_cutoff': 0.2, 'stop_bands': [], 'window': 'hann',
                      'signal_mode': 0, 'show_power': True, 'show_signal': True, 'show_legend': False},
    'Resolution enhancement': {'irf_sigma': 2.0, 'snr': 100.0},
    'X-axis alignment': {'reference_spectrum_index': 0, 'max_shift': 10.0,
                         'interpolation_method': 'cubic', 'x_range': None},
    'X-axis Unit Conversion': {'from_unit': 'nm', 'to_unit': 'cm-1'},
    'CD Unit Conversion': {'concentration_unit': 'µM', 'ellipticity_convention': 'mean_residue',
                           'output_type': 'delta_epsilon',
                           'per_spectrum': {'a': {'path_length_cm': 0.1, 'concentration_value': 5.0}}},
    'Mean-Center Spectra (Dataset)': {'include_mean_spectrum': False},
    'Combine Spectra': {'operation_type': 'average', 'new_spectrum_name': 'Avg', 'show_variance': False,
                        'source_labels': ['a', 'b', 'c']},
    'Spectral Calculator': {'formula': 'A - 0.5 * B', 'new_spectrum_name': 'diff'},
    'Interactive subtraction': {'minuend_indices': [0], 'subtrahend_index': 1, 'subtraction_factor': 0.8,
                                'stored_factors': {'a': 0.8}},
}

# Operation types described only generically (on purpose).
GENERIC_OK = {'Peak Fitting', 'Melting Curve Analysis', 'SVD Interpolation', 'Rename'}


def _record(op, params=None, n_affected=10, n_out=10):
    return {'type': op, 'parameters': params if params is not None else REAL_SETTINGS.get(op, {}),
            'affected_labels': [f's{i}' for i in range(n_affected)], 'output_spectra': [None] * n_out}


ORIG = [{'label': f's{i}', 'metadata': {'file_path': 'f.txt'}} for i in range(10)]


@pytest.mark.parametrize('op', sorted(REAL_SETTINGS))
def test_every_operation_gets_a_proper_sentence(op):
    d = mr.describe_operation(op, REAL_SETTINGS[op])
    assert not d.sentence.startswith(op + ' was applied'), d.sentence    # not the generic fallback
    assert d.sentence.endswith('.') and d.rows
    rendered = mr._resolve_refs(d.sentence, [])
    assert '{' not in rendered and 'None' not in rendered and 'np.' not in rendered


def test_every_operation_in_the_app_is_described():
    """A new operation added to the app must get a describer (or be put on
    the generic list consciously)."""
    found = set()
    for path in glob.glob(os.path.join(ROOT, 'src', 'controllers', '**', '*.py'), recursive=True):
        src = open(path, encoding='utf-8').read()
        for m in re.finditer(r"apply_operation\(\s*['\"]([^'\"]+)['\"]", src):
            found.add(m.group(1))
    assert len(found) >= 20
    missing = found - set(mr._DESCRIBERS) - GENERIC_OK
    assert not missing, f"no Methods-report description for: {sorted(missing)}"


def test_manual_baseline_counts_points_in_both_stored_shapes():
    d = mr.describe_operation('Manual baseline', REAL_SETTINGS['Manual baseline'])
    assert 'polynomial (order 3)' in d.sentence and '4–5 points per spectrum' in d.sentence


def test_report_text_numbers_references_in_order_of_first_use():
    chain = [_record('Data range'), _record('Automated Baseline'), _record('SG-smoothing'),
             _record('Automated Baseline', dict(REAL_SETTINGS['Automated Baseline'], algorithm='als')),
             _record('SG-smoothing')]
    r = mr.build_report(chain, len(chain) - 1, ORIG, '1.5.0')
    assert r.text.startswith('Spectra were processed and analysed in SpecAnalytiXBase 1.5.0 [1].')
    assert 'arPLS) [2]' in r.text and 'Savitzky–Golay filter [3, 4]' in r.text and '(ALS) [5]' in r.text
    assert r.text.count('[3, 4]') == 2                    # same references reused, not renumbered
    assert len(r.references) == 5 and 'SpecAnalytiXBase version 1.5.0' in r.references[0]
    assert r.references[1].startswith('Baek') and r.references[4].startswith('Eilers')
    assert 'λ = 10⁵' in r.text and 'at most 50 iterations' in r.text
    assert 'restricted to 220–320' in r.text and 'step of 0.5' in r.text


def test_only_steps_up_to_the_active_one_and_no_renames():
    chain = [_record('SG-smoothing'), _record('Rename', {}), _record('Normalization'), _record('SNIP Baseline')]
    r = mr.build_report(chain, 2, ORIG)
    assert [s.operation for s in r.steps] == ['SG-smoothing', 'Normalization']
    assert 'SNIP' not in r.text
    assert mr.build_report(chain, -1, ORIG).text.endswith('No processing was applied to the spectra.')


def test_a_step_on_part_of_the_spectra_says_so():
    r = mr.build_report([_record('SG-smoothing', n_affected=4)], 0, ORIG)
    assert 'polynomial order 3; applied to 4 of 10 spectra).' in r.text
    assert r.steps[0].applied_to == '4 of 10'


def test_units_follow_the_data_and_unit_conversion():
    jasco = [{'label': 'a', 'metadata': {'file_type': 'jasco_jwb', 'file_path': 'x.jwb'}}]
    r = mr.build_report([_record('Data range', n_affected=1, n_out=1)], 0, jasco)
    assert '220–320 nm' in r.text
    chain = [_record('X-axis Unit Conversion', n_affected=1, n_out=1),
             _record('Resolution enhancement', n_affected=1, n_out=1)]
    r = mr.build_report(chain, 1, jasco)
    assert 'wavelength (nm) to wavenumber (cm⁻¹)' in r.text and 'σ = 2 cm⁻¹' in r.text


def test_batch_pipeline_describes_its_steps():
    params = {'pipeline_name': 'Raman std', 'steps': [
        {'operation': 'SNIP Baseline', 'settings': REAL_SETTINGS['SNIP Baseline']},
        {'operation': 'Normalization', 'settings': REAL_SETTINGS['Normalization']}]}
    r = mr.build_report([_record('Batch Pipeline: Raman std', params)], 0, ORIG)
    assert '“Raman std”' in r.text and 'SNIP algorithm [2]' in r.text and 'unit vector length' in r.text
    assert any('step 2: Normalization' in row[3] for row in r.table_rows())


def test_unknown_or_odd_records_never_break_the_report():
    chain = [_record('Something new', {'width': np.float64(2.5), 'flags': [1, 2], 'nested': {'a': 1}}),
             _record('SG-smoothing', {'window_length': 'eleven'}),          # odd value
             _record('Automated Baseline', {'algorithm': 'brand-new'})]
    r = mr.build_report(chain, 2, ORIG)
    assert 'Something new was applied (width 2.5' in r.text
    assert len(r.steps) == 3


def test_renderings():
    r = mr.build_report([_record('SG-smoothing')], 0, ORIG, '1.5.0', {'NumPy': '2.1'})
    assert 'Calculations used NumPy 2.1.' in r.text
    assert '<table>' in r.as_html() and 'Savitzky' in r.as_html()
    assert '| Step | Operation |' in r.as_markdown()
    assert r.as_plain_text().startswith('DATA PROCESSING')
    assert mr.fmt_num(1e5) == '10⁵' and mr.fmt_num(2.5e-4) == '2.5 × 10⁻⁴' and mr.fmt_num(11) == '11'


# --- in the application -------------------------------------------------------

@pytest.fixture
def mc(tmp_path, monkeypatch):
    monkeypatch.setenv('SPECANALYTIXBASE_DATA_DIR', str(tmp_path / 'appdata'))
    for name in ('information', 'warning', 'critical'):
        monkeypatch.setattr(QMessageBox, name, staticmethod(lambda *a, **k: QMessageBox.Ok))
    monkeypatch.setattr(QMessageBox, 'question', staticmethod(lambda *a, **k: QMessageBox.Yes))
    from src.controllers.core.main_controller import MainController
    from src.views.dialogs.misc import import_dialog
    monkeypatch.setattr(import_dialog.ImportDialog, 'run', staticmethod(
        lambda view, settings, file_paths=None: (True, list(file_paths), [{} for _ in file_paths])))
    controller = MainController()
    controller.import_controller._run_import(clear_first=True, file_paths=[TBA])
    app.processEvents()
    yield controller
    controller.autosave_controller.shutdown()
    controller.view.deleteLater()


def test_menu_entry_and_dialog(mc, monkeypatch, tmp_path):
    texts = [a.text() for a in mc.view.menuFile.actions()]
    assert texts.index('Methods report for publication…') == texts.index(mc.view.actionSave.text()) + 1

    om = mc.operations_controller.operations_manager
    spectra = om.get_current_spectra()
    om.apply_operation('SG-smoothing', REAL_SETTINGS['SG-smoothing'], spectra[:5])
    report = mr.report_for_main_controller(mc)
    assert 'Savitzky–Golay' in report.text and 'applied to 5 of' in report.text
    assert 'NumPy' in report.text

    from src.views.dialogs.misc import methods_report_dialog as mod
    shown = []
    monkeypatch.setattr(mod.MethodsReportDialog, 'exec_', lambda self: shown.append(self) or 0)
    mc.show_methods_report()
    (dlg,) = shown
    dlg.text_edit.setPlainText(dlg.text_edit.toPlainText() + ' Edited by me.')
    dlg.table.item(0, 4).setText('my value')
    dlg.refs_edit.setPlainText('Baek et al. (2015), Analyst 140, 250.\n\nMy own reference.')

    dlg.include_table_cb.setChecked(True)
    dlg.include_refs_cb.setChecked(True)
    dlg.copy()
    mime = QApplication.clipboard().mimeData()
    assert 'Edited by me.' in mime.text() and 'my value' in mime.text()
    assert '<table>' in mime.html() and 'My own reference.' in mime.html()
    assert '[2]' not in mime.text().split('REFERENCES')[1]       # edited references used as written

    dlg.include_table_cb.setChecked(False)
    dlg.include_refs_cb.setChecked(False)
    dlg.copy()
    mime = QApplication.clipboard().mimeData()
    assert 'Edited by me.' in mime.text()
    assert 'my value' not in mime.text() and 'My own reference' not in mime.text()
    assert '<table>' not in mime.html()
    from src.modules.misc.session_files import Preferences
    assert Preferences().get('methods_report_include_table') is False   # remembered

    dlg.include_refs_cb.setChecked(True)
    for ext in ('.docx', '.html', '.md', '.txt'):
        target = str(tmp_path / f'report{ext}')
        monkeypatch.setattr(QFileDialog, 'getSaveFileName', staticmethod(lambda *a, **k: (target, '')))
        dlg.save()
        assert dlg.saved_path == target
        if ext == '.docx':
            import zipfile
            content = zipfile.ZipFile(target).read('word/document.xml').decode('utf-8')
        else:
            content = open(target, encoding='utf-8').read()
        assert 'Edited by me.' in content and 'Savitzky' in content and 'My own reference.' in content
        assert 'my value' not in content

    # no extension typed: the chosen format decides
    target = str(tmp_path / 'noext')
    monkeypatch.setattr(QFileDialog, 'getSaveFileName',
                        staticmethod(lambda *a, **k: (target, 'Word document (*.docx)')))
    dlg.save()
    assert dlg.saved_path == target + '.docx'


def test_help_button_opens_its_topic(mc, monkeypatch):
    from src.help import help_window
    from src.views.dialogs.misc import methods_report_dialog as mod
    opened = []
    monkeypatch.setattr(help_window, 'open_help_topic', lambda parent, key: opened.append(key))
    dlg = mod.MethodsReportDialog(mr.report_for_main_controller(mc))
    dlg.help_btn.click()
    assert opened == ['methods_report']
    assert 'methods_report' in help_window._REGISTRY
    from src.help.methods_report_help import get_methods_report_help_content
    assert 'template' in get_methods_report_help_content()


def test_docx_is_a_valid_word_document(tmp_path):
    import zipfile
    import xml.dom.minidom
    r = mr.build_report([_record('SG-smoothing'), _record('Automated Baseline')], 1, ORIG, '1.5.0')
    content = r.export_content()
    content.text += ' Special characters: <&> "quotes" λ 10⁵ µM – —\n\nSecond paragraph.'
    path = str(tmp_path / 'm.docx')
    content.save(path)
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        assert {'[Content_Types].xml', '_rels/.rels', 'word/document.xml', 'word/styles.xml'} <= names
        for name in names:
            xml.dom.minidom.parseString(z.read(name))           # every part is well-formed XML
        doc = z.read('word/document.xml').decode('utf-8')
    assert '&lt;&amp;&gt;' in doc and 'λ 10⁵ µM' in doc and '<w:tbl>' in doc
    try:
        import docx                                              # python-docx, if installed
    except ImportError:
        return
    d = docx.Document(path)
    texts = [p.text for p in d.paragraphs]
    assert texts[0] == 'Data processing' and d.paragraphs[0].style.name == 'Heading 1'
    assert any(t.startswith('Spectra were processed') for t in texts)
    assert 'Second paragraph.' in texts
    assert len(d.tables) == 1 and d.tables[0].rows[0].cells[0].text == 'Step'
    assert any(t.startswith('[2] Savitzky') for t in texts)


def test_edits_are_kept_when_the_window_is_reopened(mc, monkeypatch):
    from src.views.dialogs.misc import methods_report_dialog as mod
    shown = []
    monkeypatch.setattr(mod.MethodsReportDialog, 'exec_', lambda self: shown.append(self) or 0)
    om = mc.operations_controller.operations_manager
    om.apply_operation('SG-smoothing', REAL_SETTINGS['SG-smoothing'], om.get_current_spectra())

    mc.show_methods_report()
    first = shown[-1]
    generated_text = first.text_edit.toPlainText()
    assert first.draft_label.isHidden()
    first.text_edit.setPlainText('My own Methods text.')
    first.table.item(0, 4).setText('edited cell')
    first.refs_edit.setPlainText('My reference.')
    first.reject()                                         # closed, e.g. by accident

    mc.show_methods_report()                               # same processing -> edits back
    second = shown[-1]
    assert second.text_edit.toPlainText() == 'My own Methods text.'
    assert second.table.item(0, 4).text() == 'edited cell'
    assert second.refs_edit.toPlainText() == 'My reference.'
    assert not second.draft_label.isHidden()

    monkeypatch.setattr(QMessageBox, 'question', staticmethod(lambda *a, **k: QMessageBox.Yes))
    second.restore_generated()
    assert second.text_edit.toPlainText() == generated_text and second.draft_label.isHidden()
    second.text_edit.setPlainText('Edited again.')
    second.accept()

    # processing changed -> the user is asked
    om.apply_operation('SNIP Baseline', REAL_SETTINGS['SNIP Baseline'], om.get_current_spectra())
    answers = iter(['New draft', 'Keep my edited version'])

    def answer(box):
        wanted = next(answers)
        for b in box.buttons():
            if b.text() == wanted:
                b.click()
        return 0
    monkeypatch.setattr(QMessageBox, 'exec_', answer)
    mc.show_methods_report()
    third = shown[-1]
    assert 'SNIP' in third.text_edit.toPlainText()        # new draft
    third.text_edit.setPlainText('Edited on the new processing.')
    third.reject()
    om.apply_operation('Normalization', REAL_SETTINGS['Normalization'], om.get_current_spectra())
    mc.show_methods_report()
    assert shown[-1].text_edit.toPlainText() == 'Edited on the new processing.'   # kept on request


def test_unedited_report_leaves_no_draft(mc, monkeypatch):
    from src.views.dialogs.misc import methods_report_dialog as mod
    monkeypatch.setattr(mod.MethodsReportDialog, 'exec_', lambda self: self.reject() or 0)
    mc.show_methods_report()
    assert mc._methods_report_draft == {}
