"""
generate_kinetics_test_datasets.py

Generates 2 synthetic spectral workbooks for testing / demonstrating the
Kinetics Fitting tool: one single-wavelength single-exponential decay
demo, and one global-analysis two-step-reaction demo with two shared
rate constants and distinct Decay-Associated Spectra. Matches the
existing SYNTHETIC_TEST_DATASETS convention -- first sheet is 'Spectra'
(what gets imported), plus an Info sheet with the known ground truth.
"""
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill

SEED = 11
HEADER_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
BOLD = Font(bold=True)


def gaussian(x, center, width, amp):
    return amp * np.exp(-((x - center) ** 2) / (2 * width ** 2))


# ======================================================================
# Dataset 1 -- Single wavelength: single-exponential decay
# ======================================================================
def build_single_exp_dataset(path):
    rng = np.random.default_rng(SEED)
    wl = np.linspace(400, 700, 80)

    true_k = 0.20
    true_amp = 1.0
    true_yinf = 0.05

    times = np.round(np.concatenate([
        np.linspace(0.5, 5, 10, endpoint=False),
        np.linspace(5, 25, 10),
    ]), 2)

    labels = [f"decay_t{t:g}" for t in times]
    spectra_y = []
    band = gaussian(wl, 550, 35, 1.0)  # fixed spectral shape, amplitude decays
    for t in times:
        signal_frac = true_yinf + true_amp * np.exp(-true_k * t)
        y = signal_frac * band + rng.normal(0, 0.01, size=wl.shape)
        spectra_y.append(y)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spectra"
    header = ["Wavelength_nm"] + labels
    ws.append(header)
    for cell in ws[1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for i, w in enumerate(wl):
        row = [float(w)] + [float(s[i]) for s in spectra_y]
        ws.append(row)
    ws.column_dimensions["A"].width = 14

    ws2 = wb.create_sheet("Time_values")
    ws2.append(["Spectrum", "Time"])
    for cell in ws2[1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for lbl, t in zip(labels, times):
        ws2.append([lbl, float(t)])
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 10

    ws3 = wb.create_sheet("Info")
    lines = [
        "Story",
        "A species' absorbance band at ~550 nm decays away with a single",
        "clean first-order process as it's consumed/converted over time.",
        "Classic single-wavelength kinetics: read the signal at the band's",
        "peak and fit a single exponential decay.",
        "",
        "How to use this file",
        "1. File -> Import data -> select this workbook (Spectra sheet).",
        "2. Select all 20 imported spectra, open Visualization & Analysis ->",
        "   Kinetics Fitting.",
        "3. Leave mode on 'Single wavelength'. Time values are already",
        "   guessed correctly from the labels -- no editing needed.",
        "4. Set Extraction X to 550, Number of components to 1.",
        "5. Click Run Fit.",
        "",
        "Ground truth (for reference)",
    ]
    for text in lines:
        ws3.append([text])
    ws3["A1"].font = BOLD
    ws3.column_dimensions["A"].width = 90
    ws3.append(["True rate constant k:", true_k])
    ws3.append(["True tau (1/k):", round(1.0 / true_k, 3)])
    ws3.append(["True offset (y_inf):", true_yinf])
    ws3.append(["Expected fit result:", f"k ~ {true_k:.2f} (tau ~ {1/true_k:.1f}), R2 > 0.99"])

    wb.save(path)
    return labels, times, true_k


# ======================================================================
# Dataset 2 -- Global analysis: two-step consecutive-ish reaction
# ======================================================================
def build_global_two_step_dataset(path):
    rng = np.random.default_rng(SEED + 1)
    wl = np.linspace(350, 650, 90)

    true_k_fast = 0.5   # tau ~ 2
    true_k_slow = 0.08  # tau ~ 12.5

    das_fast = gaussian(wl, 450, 30, 1.5)   # fast-decaying band, blue side
    das_slow = -gaussian(wl, 550, 35, 1.1)  # slow, oppositely-signed band, red side
    offset_spec = 0.05 * np.ones_like(wl)

    times = np.round(np.array([0.1, 0.3, 0.6, 1, 1.5, 2, 3, 4, 6, 8, 11, 15,
                                20, 26, 33, 41, 50, 60]), 2)
    labels = [f"twostep_t{t:g}" for t in times]

    spectra_y = []
    for t in times:
        y = (das_fast * np.exp(-true_k_fast * t)
             + das_slow * np.exp(-true_k_slow * t)
             + offset_spec)
        y += rng.normal(0, 0.008, size=wl.shape)
        spectra_y.append(y)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spectra"
    header = ["Wavelength_nm"] + labels
    ws.append(header)
    for cell in ws[1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for i, w in enumerate(wl):
        row = [float(w)] + [float(s[i]) for s in spectra_y]
        ws.append(row)
    ws.column_dimensions["A"].width = 14

    ws2 = wb.create_sheet("Time_values")
    ws2.append(["Spectrum", "Time"])
    for cell in ws2[1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for lbl, t in zip(labels, times):
        ws2.append([lbl, float(t)])
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 10

    ws3 = wb.create_sheet("Info")
    lines = [
        "Story",
        "A two-step process (A -> B -> C) where the fast step (tau ~ 2) and",
        "slow step (tau ~ 12.5) each have their own, overlapping spectral",
        "signature across 350-650 nm. Reading any single wavelength mixes",
        "both processes together -- global analysis recovers both shared",
        "rate constants AND each step's own Decay-Associated Spectrum (DAS)",
        "from the whole dataset at once.",
        "",
        "How to use this file",
        "1. File -> Import data -> select this workbook (Spectra sheet).",
        "2. Select all 18 imported spectra, open Visualization & Analysis ->",
        "   Kinetics Fitting.",
        "3. Choose 'Global analysis'. Time values are already guessed",
        "   correctly from the labels -- no editing needed.",
        "4. Set Number of components to 2, click Run Fit.",
        "5. Switch the View dropdown to 'Decay-Associated Spectra' to see",
        "   each step's own spectral signature.",
        "",
        "Ground truth (for reference)",
    ]
    for text in lines:
        ws3.append([text])
    ws3["A1"].font = BOLD
    ws3.column_dimensions["A"].width = 90
    ws3.append(["True fast rate constant k:", true_k_fast])
    ws3.append(["True fast tau:", round(1.0 / true_k_fast, 3)])
    ws3.append(["True slow rate constant k:", true_k_slow])
    ws3.append(["True slow tau:", round(1.0 / true_k_slow, 3)])
    ws3.append(["Expected fit result:", f"k ~ {true_k_slow:.2f} and k ~ {true_k_fast:.2f} "
                                        f"(tau ~ {1/true_k_slow:.1f} and ~ {1/true_k_fast:.1f}), R2 > 0.99"])

    wb.save(path)
    return labels, times, true_k_fast, true_k_slow


if __name__ == "__main__":
    import os
    out_dir = os.path.dirname(os.path.abspath(__file__))
    p1 = os.path.join(out_dir, "kinetics_single_exponential_decay_demo.xlsx")
    p2 = os.path.join(out_dir, "kinetics_global_two_step_demo.xlsx")
    r1 = build_single_exp_dataset(p1)
    r2 = build_global_two_step_dataset(p2)
    print("Wrote:", p1)
    print("Wrote:", p2)
    print("Single-exp times:", list(r1[1]))
    print("Global two-step times:", list(r2[1]))
