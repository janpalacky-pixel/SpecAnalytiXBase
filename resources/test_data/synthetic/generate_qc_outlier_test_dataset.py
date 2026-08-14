"""
generate_qc_outlier_test_dataset.py

Generates 1 synthetic spectral workbook for testing / demonstrating the
QC / Outlier Detection tool: 30 "normal" batch spectra (a single Gaussian
band with realistic batch-to-batch concentration variation) plus 3
deliberately planted problem spectra -- one with an extreme concentration
(a T^2-type outlier: right shape, unusual magnitude), one with an added
spike artifact (a Q-type outlier: doesn't fit the model's shape at all),
and one with both problems at once. Matches the existing
SYNTHETIC_TEST_DATASETS convention -- first sheet is 'Spectra' (what gets
imported), plus an Info sheet with the known ground truth.
"""
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill

SEED = 23
HEADER_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
BOLD = Font(bold=True)


def gaussian(x, center, width, amp):
    return amp * np.exp(-((x - center) ** 2) / (2 * width ** 2))


def build_dataset(path):
    rng = np.random.default_rng(SEED)
    wl = np.linspace(400, 700, 150)

    labels = []
    rows = []

    n_normal = 30
    for i in range(n_normal):
        conc = rng.normal(1.0, 0.08)   # ~8% RSD batch-to-batch variation
        y = conc * gaussian(wl, 550, 40, 1.0) + rng.normal(0, 0.005, wl.shape)
        labels.append(f"batch_{i + 1:02d}")
        rows.append(y)

    # Planted outlier 1: right spectral shape, extreme concentration (T^2-type)
    y_t2 = 5.0 * gaussian(wl, 550, 40, 1.0) + rng.normal(0, 0.005, wl.shape)
    labels.append("PLANTED_T2_extreme_concentration")
    rows.append(y_t2)

    # Planted outlier 2: normal concentration, added spike artifact (Q-type)
    y_q = 1.0 * gaussian(wl, 550, 40, 1.0) + gaussian(wl, 480, 3, 0.8) + rng.normal(0, 0.005, wl.shape)
    labels.append("PLANTED_Q_spike_artifact")
    rows.append(y_q)

    # Planted outlier 3: both problems at once
    y_both = 4.0 * gaussian(wl, 550, 40, 1.0) + gaussian(wl, 620, 4, 0.6) + rng.normal(0, 0.005, wl.shape)
    labels.append("PLANTED_BOTH_conc_and_spike")
    rows.append(y_both)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spectra"
    header = ["Wavelength_nm"] + labels
    ws.append(header)
    for cell in ws[1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for i, w in enumerate(wl):
        row = [float(w)] + [float(r[i]) for r in rows]
        ws.append(row)
    ws.column_dimensions["A"].width = 14

    ws2 = wb.create_sheet("Info")
    lines = [
        "Story",
        "30 spectra of the same band, each from a different sample with",
        "normal batch-to-batch concentration variation (~8% RSD) -- this",
        "is what a clean, unremarkable batch looks like. 3 more spectra",
        "were then deliberately planted with problems:",
        "  - PLANTED_T2_extreme_concentration: the right spectral SHAPE,",
        "    but 5x the normal concentration -- an unusual but still",
        "    model-shaped spectrum (flags mainly on Hotelling T^2).",
        "  - PLANTED_Q_spike_artifact: normal concentration, but with an",
        "    extra narrow spike added at 480 nm that the batch's normal",
        "    variation doesn't include (flags on Q-residual/SPE).",
        "  - PLANTED_BOTH_conc_and_spike: an extreme concentration (4x)",
        "    AND a spike artifact together (flags on both).",
        "",
        "How to use this file",
        "1. File -> Import data -> select this workbook (Spectra sheet).",
        "2. Select all 33 imported spectra, open Visualization & Analysis ->",
        "   QC / Outlier Detection.",
        "3. Leave Auto-select components and 95% confidence at their",
        "   defaults, click Run QC Check.",
        "4. Expected result: exactly the 3 PLANTED_* spectra flagged,",
        "   at the top of the worst-first results table.",
        "",
        "Ground truth (for reference)",
    ]
    for text in lines:
        ws2.append([text])
    ws2["A1"].font = BOLD
    ws2.column_dimensions["A"].width = 90
    ws2.append(["Normal spectra:", n_normal])
    ws2.append(["Planted outliers:", 3])
    ws2.append(["Expected flagged count:", 3])

    wb.save(path)
    return labels


if __name__ == "__main__":
    import os
    out_dir = os.path.dirname(os.path.abspath(__file__))
    p = os.path.join(out_dir, "qc_outlier_planted_outliers_demo.xlsx")
    r = build_dataset(p)
    print("Wrote:", p)
    print("Labels:", r)
